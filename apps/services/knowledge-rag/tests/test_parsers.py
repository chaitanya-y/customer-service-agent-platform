import hashlib
from pathlib import Path

import pytest
from docx import Document

from knowledge_rag.models import SourceContentType
from knowledge_rag.parsers import parse_source

REPOSITORY_ROOT = Path(__file__).resolve().parents[4]
FIXTURE_DIRECTORY = (
    REPOSITORY_ROOT
    / "apps"
    / "services"
    / "control-knowledge"
    / "fixtures"
    / "source-documents"
    / "acme"
)


def test_parses_refund_policy_markdown_into_citable_sections() -> None:
    source_path = FIXTURE_DIRECTORY / "refund-policy-2026-08-01.md"

    document = parse_source(source_path)

    assert document.source_content_type == SourceContentType.MARKDOWN
    assert document.title == "Acme Refund Policy"
    assert (
        document.source_content_sha256
        == hashlib.sha256(source_path.read_bytes()).hexdigest()
    )

    damaged_items = next(
        section
        for section in document.sections
        if section.heading == "2.1 Damaged items"
    )

    assert damaged_items.heading_path == [
        "Acme Refund Policy",
        "2. Refund eligibility",
        "2.1 Damaged items",
    ]
    assert "30 calendar days" in damaged_items.text
    assert "Photo evidence is required" in damaged_items.text


def test_parses_html_and_preserves_refund_review_content() -> None:
    source_path = FIXTURE_DIRECTORY / "refund-policy-2026-08-01.html"

    document = parse_source(source_path)

    assert document.source_content_type == SourceContentType.HTML
    assert document.title == "Acme Refund Policy"

    refund_review = next(
        section
        for section in document.sections
        if section.heading == "4. Refund amount and review"
    )

    assert "USD 100.00" in refund_review.text
    assert "USD 500.00" in refund_review.text
    assert "human takeover" in refund_review.text


def test_parses_docx_and_records_its_page_location_limit() -> None:
    source_path = FIXTURE_DIRECTORY / "refund-policy-2026-08-01.docx"

    document = parse_source(source_path)

    assert document.source_content_type == SourceContentType.DOCX
    assert document.title == "Acme Refund Policy"

    damaged_items = next(
        section
        for section in document.sections
        if section.heading == "2.1 Damaged items"
    )

    assert damaged_items.heading_path == [
        "Acme Refund Policy",
        "2. Refund eligibility",
        "2.1 Damaged items",
    ]
    assert damaged_items.page_start is None
    assert damaged_items.page_end is None
    assert any(
        warning.code == "DOCX_PAGE_LOCATIONS_UNAVAILABLE"
        for warning in document.extraction_warnings
    )


def test_rejects_docx_with_table_instead_of_indexing_partial_text(
    tmp_path: Path,
) -> None:
    source_path = tmp_path / "policy-with-table.docx"
    source = Document()
    source.add_paragraph("Paragraph text alone is incomplete policy evidence.")
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Critical policy condition"
    source.save(source_path)

    with pytest.raises(ValueError, match="DOCX tables are not supported"):
        parse_source(source_path)


def test_parses_plain_text_as_one_introduction_section(tmp_path: Path) -> None:
    source_path = tmp_path / "support-note.txt"
    source_path.write_text(
        "Ask the customer for the order number before checking refund eligibility.",
        encoding="utf-8",
    )

    document = parse_source(source_path)

    assert document.source_content_type == SourceContentType.TEXT
    assert document.title == "Support Note"
    assert len(document.sections) == 1
    assert document.sections[0].heading == "Introduction"


def test_parses_pdf_with_page_accurate_evidence() -> None:
    source_path = FIXTURE_DIRECTORY / "refund-policy-2026-08-01.pdf"

    document = parse_source(source_path)

    assert document.source_content_type == SourceContentType.PDF
    assert "Acme Refund Policy" in document.title
    assert len(document.sections) == 2

    first_page = document.sections[0]

    assert first_page.heading == "Page 1"
    assert first_page.page_start == 1
    assert first_page.page_end == 1
    assert "2.1 Damaged items" in first_page.text
    assert "30 calendar days" in first_page.text

    assert any(
        warning.code == "PDF_LAYOUT_SEMANTICS_UNAVAILABLE"
        for warning in document.extraction_warnings
    )
    assert not any(
        warning.code == "PDF_PAGE_EMPTY" for warning in document.extraction_warnings
    )
